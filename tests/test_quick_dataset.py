import hashlib
import json
import tempfile
import unittest
from pathlib import Path

import numpy as np
from PIL import Image
from openpyxl import load_workbook

from foodquality8.quick_dataset import prepare_google, filter_image, suggest_stage
from foodquality8.model_data import check_release
from image_search import export_excel


class QuickDatasetTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root=Path(self.temp.name)
        self.folder=self.root/'food15_workspace/image_downloads/beef'
        self.folder.mkdir(parents=True)
        self.rows=[]
        for i in range(9):
            path=self.folder/f'{i}.jpg'
            Image.fromarray(np.random.default_rng(i).integers(0,256,(300,300,3),dtype=np.uint8)).save(path)
            self.rows.append(dict(food_code=f'beef_{i}',img_path=str(path),sha256=hashlib.sha256(path.read_bytes()).hexdigest(),
                search_stage=('fresh','near_spoiled','spoiled')[i//3],width=300,height=300,source_url=f'https://example.org/{i}.jpg'))
        self.save()

    def save(self):
        (self.folder/'manifest.jsonl').write_text('\n'.join(json.dumps(r) for r in self.rows),encoding='utf8')

    def test_splits_are_valid_but_suggested_data_cannot_be_used_as_verified(self):
        release=prepare_google(self.root,'beef')
        info=check_release(release,'classifier',experimental=True)
        self.assertEqual(info['label_mode'],'google_query_suggested')
        self.assertEqual(sum(info['counts']['train'].values()),2)
        with self.assertRaises(ValueError):
            check_release(release,'classifier')
        with self.assertRaises(ValueError):
            check_release(release,'detector',experimental=True)
        self.assertFalse((self.root/'food15_workspace/catalog.db').exists())

    def test_excel_rejection_is_respected_and_report_written(self):
        excel=self.root/'review.xlsx'
        export_excel('beef',self.rows,excel)
        wb=load_workbook(excel)
        wb['beef']['E2']='Loai bo'
        wb.save(excel)
        wb.close()
        with self.assertRaisesRegex(ValueError,'2/3'):
            prepare_google(self.root,'beef',excel)
        report=json.loads((self.folder/'filter-report.json').read_text(encoding='utf8'))
        self.assertEqual(report['counts']['fresh'],2)
        wb=load_workbook(excel)
        self.assertEqual(wb['beef']['E2'].value,'Loai bo')
        self.assertEqual(len(wb['beef']._images),9)
        wb.close()

    def test_conflicting_duplicate_is_excluded(self):
        self.rows[6].update(img_path=self.rows[0]['img_path'],sha256=self.rows[0]['sha256'])
        self.save()
        with self.assertRaises(ValueError):
            prepare_google(self.root,'beef')
        report=json.loads((self.folder/'filter-report.json').read_text(encoding='utf8'))
        self.assertEqual(report['counts']['fresh'],2)
        self.assertEqual(report['counts']['spoiled'],2)

    def test_technical_filter_and_legacy_ambiguity(self):
        path=self.folder/'blank.jpg'
        Image.new('RGB',(300,300),'white').save(path)
        self.assertIsNone(filter_image(path)[1])
        self.assertIsNone(filter_image(self.rows[0]['img_path'],'https://www.shutterstock.com/a.jpg')[1])
        self.assertIsNone(suggest_stage({'query':'beef semi fresh food'}))
        self.assertIsNone(suggest_stage({'search_stage':'less_fresh'}))
