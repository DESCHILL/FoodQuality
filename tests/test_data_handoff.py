import hashlib
import json
from pathlib import Path
import shutil
import tempfile
import unittest

import numpy as np
from PIL import Image
from openpyxl import load_workbook

from foodquality8.dataset.exchange import export_excel
from foodquality8.dataset.auto_workflow import import_local, food_id
from foodquality8.dataset.quick_dataset import prepare_google


class DataHandoffTests(unittest.TestCase):
    def test_vietnamese_lamb_name(self):
        self.assertEqual(food_id('Thịt cừu'), 'mutton')

    def test_relative_originals_survive_transfer_and_keep_source(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            original = root/'collector'
            original.mkdir()
            rows = []
            for index in range(12):
                path = original/'images'/f'{index}.png'
                path.parent.mkdir(exist_ok=True)
                Image.fromarray(np.random.default_rng(index).integers(0,256,(256,256,3),dtype=np.uint8)).save(path)
                rows.append(dict(food_code=str(index),img_path=str(path),width=256,height=256,
                    sha256=hashlib.sha256(path.read_bytes()).hexdigest(),
                    search_stage='fresh' if index<6 else 'spoiled',source_url=f'https://example.org/{index}.png',
                    query='tomato fresh' if index<6 else 'tomato spoiled'))
            excel = original/'ThucPham.xlsx'
            export_excel('Cà chua', rows, excel)
            workbook = load_workbook(excel)
            for index in range(12):
                workbook.active.cell(index+2,9).value=f'images/{index}.png'
            workbook.save(excel)
            workbook.close()
            transferred = root/'received'
            shutil.copytree(original,transferred)
            imported = import_local(root/'model', transferred/'ThucPham.xlsx','tomato')
            records = [json.loads(line) for line in (imported/'manifest.jsonl').read_text(encoding='utf8').splitlines()]
            self.assertEqual({r['sha256'] for r in records}, {r['sha256'] for r in rows})
            self.assertTrue(all(r['width']==256 and r['source_url'].startswith('https://example.org/') for r in records))
            release = Path(prepare_google(root/'model','tomato',source_folder=imported))
            self.assertTrue((release/'manifest.json').is_file())
            for split in ('train','val','test'):
                for condition in ('fresh','spoiled'):
                    self.assertTrue(list((release/'classification'/split/f'tomato_{condition}').glob('*.jpg')))


if __name__ == '__main__':
    unittest.main()
