import json
import tempfile
import unittest
import io
import zipfile
from unittest.mock import patch
from unittest.mock import MagicMock
from pathlib import Path

import numpy as np
from PIL import Image
from openpyxl import load_workbook

from foodquality8 import Dataset


class PipelineTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.base = Path(self.temp.name)
        self.d = Dataset(self.base / "workspace")
        self.source = dict(id="own", license_status="Approved", license_id="OWN", license_url="https://example.org/rights",
                           source_url="https://example.org/own", checked_at="2026-10-16", mapping={"FreshApple": ["apple", "fresh"]})
        self.configure()
        self.raw = self.base / "input/FreshApple"
        self.raw.mkdir(parents=True)
        rng = np.random.default_rng(123)
        for i in range(20):
            Image.fromarray(rng.integers(0, 256, (64, 64, 3), dtype=np.uint8)).save(self.raw / f"{i}.png")

    def configure(self):
        p = self.base / "sources.json"
        p.write_text(json.dumps({"sources": [self.source]}))
        self.d.configure(p)

    def tearDown(self):
        self.d.close()
        self.temp.cleanup()

    def approve(self, mutate=None):
        path = self.d.review_export()
        wb = load_workbook(path)
        ws = wb["Images"]
        h = {c.value: c.column for c in ws[1]}
        for i in range(2, ws.max_row + 1):
            if not ws.cell(i, h["error"]).value:
                ws.cell(i, h["review_status"], "Approved")
        if mutate:
            mutate(ws, h)
        wb.save(path)
        wb.close()
        return self.d.review_import(path)

    def test_end_to_end_and_idempotence(self):
        self.assertEqual(self.d.ingest("own", self.raw.parent), 20)
        self.assertEqual(self.d.ingest("own", self.raw.parent), 0)
        with self.assertRaises(ValueError): self.d.export("premature")
        self.approve()
        sizes = self.d.split()
        self.assertEqual(sum(sizes.values()), 20)
        self.assertTrue(all(sizes.values()), sizes)
        first = {r["image_id"]: r["split"] for r in self.d.rows()}
        self.d.split()
        self.assertEqual(first, {r["image_id"]: r["split"] for r in self.d.rows()})
        out = Path(self.d.export("v1"))
        self.assertEqual(len(list((out / "classification").rglob("*.jpg"))), 20)
        self.assertEqual(len(list((out / "detection").rglob("*.jpg"))), 0)
        with self.assertRaises(ValueError): self.d.export("v1")
        with self.assertRaises(ValueError): self.d.export("public", public=True)

    def test_group_and_exact_duplicate(self):
        (self.raw / "copy.png").write_bytes((self.raw / "0.png").read_bytes())
        self.d.ingest("own", self.raw.parent)
        self.approve(lambda ws, h: [ws.cell(i, h["capture_group"], "session-1") for i in (2, 3, 4)])
        self.d.split()
        rows = self.d.rows()
        self.assertEqual(sum(bool(r["duplicate_of"]) for r in rows), 1)
        self.assertEqual(len({r["split"] for r in rows if r["capture_group"] == "session-1" and self.d.eligible(r)}), 1)

    def test_pending_and_corrupt(self):
        self.source["license_status"] = "Pending"
        self.configure()
        (self.raw / "bad.jpg").write_text("bad image")
        self.d.ingest("own", self.raw.parent)
        self.approve()
        self.assertEqual(sum(self.d.split().values()), 0)
        self.assertEqual(sum(bool(r["error"]) for r in self.d.rows()), 1)
        with self.assertRaises(ValueError): self.d.download("own")

    def test_invalid_review_is_atomic(self):
        self.d.ingest("own", self.raw.parent)
        with self.assertRaises(ValueError):
            self.approve(lambda ws, h: ws.cell(3, h["condition"], "cooked"))
        self.assertTrue(all(r["review_status"] == "Pending" for r in self.d.rows()))

    def test_detection_and_tampering(self):
        self.d.ingest("own", self.raw.parent)
        def boxes(ws, h):
            ws.cell(2, h["bbox_json"], json.dumps([dict(food_name="apple", x=.5, y=.5, w=.8, h=.8)]))
            ws.cell(2, h["bbox_status"], "Approved")
        self.approve(boxes)
        self.d.split()
        out = Path(self.d.export("bbox"))
        self.assertEqual(len(list((out / "detection/labels").rglob("*.txt"))), 1)
        detection=next((out/'detection/images').rglob('*.jpg'))
        classification=next((out/'classification').rglob(detection.name))
        with Image.open(detection) as full, Image.open(classification) as crop:
            self.assertLess(crop.width,full.width)
            self.assertLess(crop.height,full.height)
        r = self.d.rows()[0]
        (self.d.root / r["raw_path"]).write_bytes(b"tampered")
        with self.assertRaises(ValueError): self.d.export("tampered")

    def test_conflicting_labels(self):
        self.source["mapping"]["RottenApple"] = ["apple", "visibly_spoiled"]
        self.configure()
        rotten = self.raw.parent / "RottenApple"
        rotten.mkdir()
        (rotten / "copy.png").write_bytes((self.raw / "0.png").read_bytes())
        self.d.ingest("own", self.raw.parent)
        self.approve()
        conflicts = [r for r in self.d.rows() if "Conflicting" in r["notes"]]
        self.assertEqual(len(conflicts), 2)
        self.assertTrue(all(r["review_status"] == "Pending" for r in conflicts))

    def test_explicit_subset_keeps_project_taxonomy(self):
        self.d.ingest('own',self.raw.parent)
        self.approve()
        self.d.split()
        original=list(self.d.classes)
        out=Path(self.d.export('pilot',selected_classes=['apple_fresh']))
        manifest=json.loads((out/'manifest.json').read_text())
        self.assertEqual(manifest['scope'],'explicit_subset')
        self.assertEqual(manifest['classes'],['apple_fresh'])
        self.assertEqual(self.d.classes,original)
        self.assertIn('names: ["apple"]',(out/'detection/data.yaml').read_text())
        with self.assertRaises(ValueError):
            self.d.export('bad',selected_classes=['invented_class'])

    def test_box_validation(self):
        for values in ((.5,.5,2,1), (.5,.5,-1,1), (float("nan"),.5,1,1)):
            with self.assertRaises(ValueError):
                self.d.validate_boxes([dict(zip(("food_name","x","y","w","h"), ("apple", *values)))])

    def test_zip_download_and_traversal(self):
        self.source['download_url'] = 'https://example.org/data.zip'
        self.configure()
        archive = io.BytesIO()
        with zipfile.ZipFile(archive, 'w') as z:
            z.writestr('../escape.txt', 'bad')
        response = MagicMock()
        response.__enter__.return_value = response
        response.iter_content.return_value = [archive.getvalue()]
        with patch('requests.get', return_value=response):
            with self.assertRaises(ValueError): self.d.download('own')
        self.assertFalse((self.d.root / 'data/raw/own/download').exists())
        archive = io.BytesIO()
        with zipfile.ZipFile(archive, 'w') as z:
            z.writestr('FreshApple/0.png', (self.raw / '0.png').read_bytes())
        response.iter_content.return_value = [archive.getvalue()]
        with patch('requests.get', return_value=response):
            folder = self.d.download('own')
        self.assertEqual(self.d.ingest('own', folder), 1)

    def test_near_duplicates_stay_together(self):
        with Image.open(self.raw / '0.png') as im:
            im.resize((128,128)).save(self.raw / 'resized.png')
        self.d.ingest('own', self.raw.parent)
        self.approve()
        self.d.split()
        rows = [r for r in self.d.rows() if r['original_path'] in ('FreshApple/0.png', 'FreshApple/resized.png')]
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[0]['group_id'], rows[1]['group_id'])
        self.assertEqual(rows[0]['split'], rows[1]['split'])

    def test_custom_taxonomy_requires_evidence_and_complete_export(self):
        taxonomy = self.base / 'taxonomy.json'
        taxonomy.write_text(json.dumps({'foods': {'apple':'fruit','guava':'fruit'},
                                       'conditions':['fresh','less_fresh','near_spoilage','slightly_spoiled','severely_spoiled']}))
        self.assertEqual(self.d.set_taxonomy(taxonomy), 10)
        self.d.ingest('own', self.raw.parent)
        with self.assertRaisesRegex(ValueError, 'reviewer and label_evidence'):
            self.approve()
        def evidence(ws, h):
            for i in range(2, ws.max_row+1):
                ws.cell(i,h['reviewer'],'test reviewer')
                ws.cell(i,h['label_evidence'],'Synthetic test fixture only')
        self.approve(evidence)
        self.d.split()
        with self.assertRaisesRegex(ValueError, 'Incomplete target dataset'):
            self.d.export('incomplete')
        coverage = self.d.coverage()
        self.assertEqual(coverage['approved_classes'], 1)
        self.assertFalse(coverage['ready_for_training'])
        self.d.close()
        self.d = Dataset(self.base / 'workspace')
        self.assertEqual(len(self.d.classes),10)

    def test_known_food_unknown_grade_stays_pending(self):
        self.source['mapping']['FreshApple']=['apple','unknown']
        self.configure()
        self.d.ingest('own',self.raw.parent)
        self.assertTrue(all(r['food_name']=='apple' and r['condition']=='unknown' and r['original_label']=='FreshApple' for r in self.d.rows()))
        self.assertEqual(sum(self.d.split().values()),0)


if __name__ == "__main__":
    unittest.main()
