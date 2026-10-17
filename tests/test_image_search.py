import json
import tempfile
import threading
import unittest
import zipfile
from pathlib import Path
from unittest.mock import MagicMock, patch

from PIL import Image
from openpyxl import load_workbook

from image_search import build_search_plan, export_excel, image_candidates, search_and_export


class ImageSearchTests(unittest.TestCase):
    def test_plan_covers_all_stages_and_translates_vietnamese_food(self):
        plan = build_search_plan('  Thịt bò  ')
        self.assertEqual(len(plan), 4)
        self.assertEqual(len({p['stage'] for p in plan[:2]}), 2)
        self.assertEqual(plan[0]['query'], 'Thịt bò tươi thực phẩm')
        self.assertEqual(plan[2]['query'], 'beef fresh food')
        self.assertEqual(plan[-1]['query'], 'beef spoiled food')
        self.assertIn('rau cải', build_search_plan('rau cải')[3]['query'])

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.img = self.root / 'fixture.jpg'
        import numpy as np
        Image.fromarray(np.random.default_rng(42).integers(0,256,(300,400,3),dtype=np.uint8)).save(self.img)
        self.row = dict(food_code='apple_1', img_path=str(self.img), sha256='abc',
                        pixel_hash='pixels', width=400, height=300, source_url='https://example.org/a.jpg')

    def test_embedded_images_survive_save_reload_and_keep_review(self):
        xlsx = self.root / 'result.xlsx'
        export_excel('apple', [self.row], xlsx)
        wb = load_workbook(xlsx)
        wb['apple']['D2'] = 'Tuoi'
        wb['apple']['E2'] = 'Da duyet'
        wb.save(xlsx)
        wb.close()
        export_excel('banana', [self.row], xlsx)
        export_excel('apple', [self.row], xlsx)
        with zipfile.ZipFile(xlsx) as z:
            self.assertEqual(len([n for n in z.namelist() if n.startswith('xl/media/')]), 2)
        wb = load_workbook(xlsx)
        self.assertEqual(wb['apple']['D2'].value, 'Tuoi')
        self.assertEqual(wb['apple']['E2'].value, 'Da duyet')
        self.assertEqual(len(wb['apple']._images), 1)
        wb.close()

    def test_stop_exports_partial_download_and_resume_target_needs_no_browser(self):
        event = threading.Event()
        page = MagicMock()
        page.url = 'https://www.google.com/search?q=apple'
        page.locator.return_value.inner_text.return_value = 'Apple images'
        page.content.return_value = '[["https://example.org/a.jpg",300,400]]'
        context = MagicMock()
        context.pages = [page]
        driver = MagicMock()
        driver.chromium.launch_persistent_context.return_value = context
        manager = MagicMock()
        manager.__enter__.return_value = driver
        def log(message):
            if message.startswith('[Anh '):
                event.set()
        xlsx = self.root / 'out.xlsx'
        with patch('playwright.sync_api.sync_playwright', return_value=manager), \
             patch('image_search._download', return_value=self.row.copy()):
            result = search_and_export('apple', 5, self.root / 'images', xlsx, event, log)
        self.assertEqual(Path(result), xlsx)
        wb = load_workbook(xlsx)
        self.assertEqual(wb['apple'].max_row, 2)
        self.assertEqual(len(wb['apple']._images), 1)
        self.assertEqual(wb['apple']['D2'].value, 'Tuoi')
        self.assertEqual(wb['apple']['E2'].value, 'Goi y tu truy van')
        wb.close()
        context.close.assert_called_once()
        with patch('playwright.sync_api.sync_playwright', side_effect=AssertionError('Should not open')):
            search_and_export('apple', 1, self.root / 'images', xlsx, threading.Event(), lambda s: None)

    def test_original_urls_not_google_thumbnails(self):
        content = '[["https://site.org/a.jpg?a=1\\u0026b=2",300,400],["https://encrypted-tbn0.gstatic.com/x",10,20]]'
        self.assertEqual(image_candidates(content), ['https://site.org/a.jpg?a=1&b=2'])

    def test_empty_stop_does_not_overwrite_existing_workbook(self):
        target = self.root / 'existing.xlsx'
        export_excel('banana', [self.row], target)
        original = target.read_bytes()
        event = threading.Event()
        event.set()
        driver = MagicMock()
        with patch('playwright.sync_api.sync_playwright', return_value=driver):
            with self.assertRaises(RuntimeError):
                search_and_export('apple', 0, self.root / 'images', target, event, lambda s: None)
        self.assertEqual(target.read_bytes(), original)


if __name__ == '__main__':
    unittest.main()
