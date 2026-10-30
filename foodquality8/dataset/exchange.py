"""Offline image filtering and spreadsheet interchange; no crawler or model dependencies."""
import hashlib, io, re
from pathlib import Path
from urllib.parse import urlparse
import numpy as np
from PIL import Image, ImageOps
from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
STAGES={"fresh":"Tươi","spoiled":"Hỏng"}


def phash(image):
    a = np.asarray(image.convert("L").resize((32, 32)), dtype=float)
    basis = np.cos(np.pi / 32 * (np.arange(32) + .5)[None, :] * np.arange(8)[:, None])
    low = (basis @ a @ basis.T).flatten()
    bits = low > np.median(low[1:])
    return f"{sum(int(b) << i for i, b in enumerate(bits)):016x}"

def suggest_stage(row):
    stage=row.get('search_stage')
    if stage in STAGES:
        return stage
    # Preserve ambiguous older labels rather than inventing equivalence.
    if stage:
        return None
    query=row.get('query','').lower()
    if any(word in query for word in ('semi','slightly','ít tươi','hơi hỏng','rất hỏng','sắp hỏng','early spoilage')):
        return None
    for key,words in [('near_spoiled',['sắp hỏng','early spoilage']),('spoiled',['rotten','spoiled','thối','mốc',' hỏng']),('fresh',[' fresh',' tươi'])]:
        if any(word in query for word in words):
            return key
    return None

def filter_image(path, url=''):
    host=(urlparse(url).hostname or '').lower()
    if any(host==d or host.endswith('.'+d) for d in ('shutterstock.com','istockphoto.com','gettyimages.com','pngtree.com','freepik.com','dreamstime.com')):
        return 'Loại: nguồn ảnh stock/minh họa',None
    try:
        with Image.open(path) as source:
            if source.width*source.height>40_000_000:
                return 'Loại: ảnh quá lớn',None
            im=ImageOps.exif_transpose(source).convert('RGB')
            if min(im.size)<200:
                return 'Loại: ảnh nhỏ hơn 200 px',None
            gray=np.asarray(im.convert('L').resize((256,256)),dtype=float)
            lap=-4*gray[1:-1,1:-1]+gray[:-2,1:-1]+gray[2:,1:-1]+gray[1:-1,:-2]+gray[1:-1,2:]
            if gray.std()<5 or lap.var()<8:
                return 'Loại: ảnh gần trống hoặc quá mờ',None
            return 'Đạt lọc kỹ thuật; chưa kiểm chứng nội dung',dict(phash=phash(im),size=im.size,
                pixel_hash=hashlib.sha256(str(im.size).encode()+im.tobytes()).hexdigest())
    except (OSError,ValueError) as exc:
        return 'Loại: không đọc được ảnh',None

def food_key(name):
    slug = re.sub(r'[^\w-]+', '_', name.strip().lower(), flags=re.UNICODE)[:60]
    return slug or 'food'

def sheet_name(name):
    safe = re.sub(r'[\\/*?\[\]:]', '_', name).strip("'") or 'Food'
    return safe if len(safe) <= 31 else safe[:24] + '_' + hashlib.sha256(name.encode()).hexdigest()[:6]

def _text(cell, value):
    # Untrusted search terms/URLs must never become spreadsheet formulas.
    cell.value = str(value or '')
    cell.data_type = 's'

def export_excel(food_name, rows, excel_path, log_cb=print):
    """Embed PNG bytes in XLSX; preserve other foods and manually reviewed labels."""
    excel_path = Path(excel_path)
    excel_path.parent.mkdir(parents=True, exist_ok=True)
    wb = load_workbook(excel_path) if excel_path.exists() else Workbook()
    title = sheet_name(food_name)
    reviews = {}
    if title in wb.sheetnames:
        old = wb[title]
        if old.cell(1, 4).value == 'Chat luong thuc pham (duyet)':
            reviews = {r[0]: (r[3], r[4]) for r in old.iter_rows(min_row=2, values_only=True) if r[0]}
        del wb[title]
    if 'Sheet' in wb.sheetnames and wb['Sheet'].max_row == 1 and wb['Sheet']['A1'].value is None:
        del wb['Sheet']
    ws = wb.create_sheet(title)
    headers = ['Ma thuc pham', 'Ten thuc pham', 'Anh thuc pham',
               'Chat luong thuc pham (duyet)', 'Trang thai duyet', 'Tu khoa tim (khong phai nhan)',
               'URL anh goc', 'Trang tim Google', 'Tep anh tren may', 'Do phan giai',
               'Thoi gian tai UTC', 'Quyen su dung', 'SHA256', 'Nhãn gợi ý từ truy vấn', 'Kết quả lọc kỹ thuật']
    widths = [24, 20, 24, 25, 20, 30, 45, 40, 40, 16, 24, 25, 32, 28, 35]
    for col, (header, width) in enumerate(zip(headers, widths), 1):
        cell = ws.cell(1, col, header)
        cell.font = Font(color='FFFFFF', bold=True)
        cell.fill = PatternFill('solid', fgColor='174A45')
        cell.alignment = Alignment(wrap_text=True, vertical='center')
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.row_dimensions[1].height = 34
    for idx, row in enumerate(rows, 2):
        label, status = reviews.get(row['food_code'], ('Chua xac dinh', 'Cho duyet'))
        if status in ('Cho duyet','Goi y tu truy van') and label=='Chua xac dinh':
            label={'fresh':'Tuoi','near_spoiled':'Sap hong','spoiled':'Hong'}.get(row.get('search_stage'),'Chua xac dinh')
            if label!='Chua xac dinh':
                status='Goi y tu truy van'
        values = [row['food_code'], food_name, '', label, status, row.get('query', ''),
                  row.get('source_url', ''), row.get('google_url', ''), str(row['img_path']),
                  f"{row['width']}x{row['height']}", row.get('downloaded_at', ''),
                  'Chua xac minh giay phep', row['sha256'], row.get('search_stage_name', 'Chua ghi nhan'), row.get('filter_result','Chưa kiểm tra lại')]
        for col, value in enumerate(values, 1):
            _text(ws.cell(idx, col), value)
            ws.cell(idx, col).alignment = Alignment(vertical='center', wrap_text=True)
        for col in (4, 5):
            ws.cell(idx, col).fill = PatternFill('solid', fgColor='FFF2CC')
        try:
            with Image.open(row['img_path']) as original:
                preview = ImageOps.exif_transpose(original).convert('RGB')
                preview.thumbnail((160, 110))
                buffer = io.BytesIO()
                preview.save(buffer, format='PNG')
                buffer.seek(0)
                picture = XLImage(buffer)
                ws.add_image(picture, f'C{idx}')
        except (OSError,ValueError):
            _text(ws.cell(idx,3),'Không đọc được ảnh — loại khỏi huấn luyện')
        ws.row_dimensions[idx].height = 85
    for col, options in [('D', 'Chua xac dinh,Tuoi,Hong'),
                         ('E', 'Goi y tu truy van,Cho duyet,Da duyet,Loai bo')]:
        validation = DataValidation(type='list', formula1=f'"{options}"')
        ws.add_data_validation(validation)
        validation.add(f'{col}2:{col}{max(2, len(rows)+1)}')
    ws.freeze_panes = 'D2'
    ws.auto_filter.ref = ws.dimensions
    temporary = excel_path.with_name(excel_path.stem + '.writing.xlsx')
    wb.save(temporary)
    wb.close()
    try:
        temporary.replace(excel_path)
        return excel_path
    except PermissionError:
        fallback = excel_path.with_name(excel_path.stem + '-moi.xlsx')
        temporary.replace(fallback)
        log_cb(f'[Excel dang mo] Da luu ban moi: {fallback}')
        return fallback
