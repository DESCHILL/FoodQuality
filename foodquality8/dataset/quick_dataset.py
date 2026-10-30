"""Explicit experimental classification data from query suggestions, never verified labels."""
import csv
import hashlib
import json
import shutil
import tempfile
import unicodedata
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path
from urllib.parse import urlparse

import numpy as np
from PIL import Image, ImageOps
from openpyxl import load_workbook
from foodquality8.dataset.core import phash
from foodquality8.dataset.google_review import FOOD_ENGLISH
from foodquality8.dataset.model_data import write_json

STAGES={'fresh':'Tươi','spoiled':'Hỏng'}


def suggest_stage(row):
    stage=row.get('search_stage')
    if stage in STAGES:
        return stage
    # Preserve ambiguous older labels rather than inventing equivalence.
    if stage:
        return None
    query=row.get('query','').lower()
    if any(word in query for word in ('semi','slightly','ít tươi','hơi hỏng','rất hỏng','sắp hỏng','early spoilage')):
        return None
    for key,words in [('near_spoiled',['sắp hỏng','early spoilage']),('spoiled',['rotten','spoiled','thối','mốc',' hỏng']),('fresh',[' fresh',' tươi'])]:
        if any(word in query for word in words):
            return key
    return None


def filter_image(path, url=''):
    host=(urlparse(url).hostname or '').lower()
    if any(host==d or host.endswith('.'+d) for d in ('shutterstock.com','istockphoto.com','gettyimages.com','pngtree.com','freepik.com','dreamstime.com')):
        return 'Loại: nguồn ảnh stock/minh họa',None
    try:
        with Image.open(path) as source:
            if source.width*source.height>40_000_000:
                return 'Loại: ảnh quá lớn',None
            im=ImageOps.exif_transpose(source).convert('RGB')
            if min(im.size)<200:
                return 'Loại: ảnh nhỏ hơn 200 px',None
            gray=np.asarray(im.convert('L').resize((256,256)),dtype=float)
            lap=-4*gray[1:-1,1:-1]+gray[:-2,1:-1]+gray[2:,1:-1]+gray[1:-1,:-2]+gray[1:-1,2:]
            if gray.std()<5 or lap.var()<8:
                return 'Loại: ảnh gần trống hoặc quá mờ',None
            return 'Đạt lọc kỹ thuật; chưa kiểm chứng nội dung',dict(phash=phash(im),size=im.size,
                pixel_hash=hashlib.sha256(str(im.size).encode()+im.tobytes()).hexdigest())
    except (OSError,ValueError) as exc:
        return 'Loại: không đọc được ảnh',None


def prepare_google(project, food_name, excel_path=None, source_folder=None, required_stages=None):
    from foodquality8.dataset.exchange import food_key, sheet_name, export_excel
    project=Path(project)
    folder=Path(source_folder) if source_folder else project/'food15_workspace/image_downloads'/food_key(food_name)
    manifest=folder/'manifest.jsonl'
    if not manifest.exists():
        raise ValueError('Chưa có ảnh của thực phẩm này. Hãy tìm ảnh trước.')
    rows=[json.loads(line) for line in manifest.read_text(encoding='utf8').splitlines() if line.strip()]
    stages=required_stages or list(STAGES)
    corrections=folder/'corrections.json'
    corrections=json.loads(corrections.read_text(encoding='utf8')) if corrections.exists() else {}
    overrides={}
    if excel_path and Path(excel_path).is_file():
        wb=load_workbook(excel_path,read_only=True,data_only=False)
        try:
            if sheet_name(food_name) in wb.sheetnames:
                for row in wb[sheet_name(food_name)].iter_rows(min_row=2,values_only=True):
                    if row and row[0]:
                        overrides[row[0]]=(row[3],row[4])
        finally:
            wb.close()
    normalized=''.join(c for c in unicodedata.normalize('NFD',food_name.lower().replace('đ','d')) if unicodedata.category(c)!='Mn')
    food=FOOD_ENGLISH.get(normalized,normalized).replace(' ','_')
    if not food or not all(c.isascii() and (c.isalnum() or c=='_') for c in food):
        food='food_'+hashlib.sha256(food_name.encode()).hexdigest()[:8]
    candidates=[]
    audit=[]
    for row in rows:
        record=dict(code=row['food_code'],result='',stage='')
        path=Path(row['img_path']).resolve()
        if not path.is_relative_to(folder.resolve()) or not path.is_file():
            record['result']='Loại: thiếu tệp hoặc đường dẫn sai'
            audit.append(record)
            continue
        if hashlib.sha256(path.read_bytes()).hexdigest()!=row['sha256']:
            record['result']='Loại: tệp đã bị thay đổi'
            audit.append(record)
            continue
        grade,status=overrides.get(row['food_code'],('', ''))
        reason,quality=filter_image(path,row.get('source_url',''))
        stage=suggest_stage(row)
        correction=corrections.get(row['food_code'],{})
        if correction.get('excluded'):
            reason,quality='Loại: người dùng bỏ ảnh',None
        if correction.get('stage') in STAGES:
            stage=correction['stage']
        if status=='Loai bo':
            reason,quality='Loại: đã đánh dấu trong Excel',None
        if status=='Da duyet':
            mapping={'Tuoi':'fresh','Tươi':'fresh','Sap hong':'near_spoiled','Sắp hỏng':'near_spoiled','Hong':'spoiled','Hỏng':'spoiled'}
            stage=mapping.get(grade)
        if correction.get('stage') in STAGES:
            stage=correction['stage']
        if quality and stage not in STAGES:
            reason,quality='Loại: nhãn chưa thuộc Tươi/Hỏng; không tự gộp mức trung gian',None
            stage=None
        row['filter_result']=reason
        if stage:
            row['search_stage_name']=STAGES[stage]
        record.update(result=reason,stage=stage or '')
        audit.append(record)
        if quality:
            candidates.append(dict(row=row,quality=quality,stage=stage,audit=record))
    # Connected near-duplicate groups never cross train/val/test. Conflicting groups are excluded.
    parents=list(range(len(candidates)))
    def find(i):
        while parents[i]!=i:
            parents[i]=parents[parents[i]]
            i=parents[i]
        return i
    for i,a in enumerate(candidates):
        for j,b in enumerate(candidates[:i]):
            if (int(a['quality']['phash'],16)^int(b['quality']['phash'],16)).bit_count()<=6:
                parents[find(i)]=find(j)
    groups=defaultdict(list)
    for i,candidate in enumerate(candidates):
        groups[find(i)].append(candidate)
    accepted=defaultdict(list)
    for group in groups.values():
        if len({r['stage'] for r in group})>1:
            reason='Loại: nhóm ảnh gần trùng có nhãn mâu thuẫn'
            keep=[]
        else:
            keep=[group[0]]
            reason='Loại: ảnh trùng/gần trùng'
            accepted[group[0]['stage']].append(group[0])
        for row in group:
            if row not in keep:
                row['row']['filter_result']=reason
                row['audit']['result']=reason
    report=dict(food=food_name,label_mode='google_query_suggested',counts={s:len(accepted[s]) for s in STAGES},
                warning='Lọc kỹ thuật không xác nhận đúng thực phẩm hoặc mức độ hỏng. Nhãn và số đo chỉ để thử nghiệm.',images=audit)
    write_json(folder/'filter-report.json',report)
    if excel_path:
        valid=[r for r in rows if Path(r['img_path']).is_file() and r.get('filter_result')]
        if valid:
            try:
                report['excel_path']=str(export_excel(food_name,valid,Path(excel_path)))
            except (OSError,ValueError) as exc:
                report['excel_error']=str(exc)
    write_json(folder/'filter-report.json',report)
    missing=[f'{STAGES[s]}: {len(accepted[s])}/3 nhóm tối thiểu' for s in stages if len(accepted[s])<3]
    if missing:
        raise ValueError('Chưa đủ ảnh sau lọc: '+ '; '.join(missing)+'. Tìm thêm ảnh Tươi/Hỏng. Chi tiết: '+str(folder/'filter-report.json'))
    releases=project/'food15_workspace/experimental_releases'
    releases.mkdir(parents=True,exist_ok=True)
    dest=releases/('google-'+food+'-'+datetime.now().strftime('%Y%m%d-%H%M%S-%f'))
    metadata=[]
    sources={}
    with tempfile.TemporaryDirectory(dir=releases) as temp:
        stage_path=Path(temp)/'release'
        stage_path.mkdir()
        for condition,items in accepted.items():
            if condition not in stages:
                continue
            items.sort(key=lambda item:item['row']['sha256'])
            holdout=max(1,round(len(items)*.15))
            for i,item in enumerate(items):
                split='test' if i<holdout else 'val' if i<holdout*2 else 'train'
                row=item['row']
                iid=row['sha256'][:24]
                target=stage_path/'classification'/split/(food+'_'+condition)/(iid+'.jpg')
                target.parent.mkdir(parents=True,exist_ok=True)
                with Image.open(row['img_path']) as image:
                    ImageOps.exif_transpose(image).convert('RGB').save(target,quality=95)
                sources[iid]=dict(source_url=row.get('source_url',''),license_status='Pending',query=row.get('query',''))
                metadata.append(dict(image_id=iid,food_name=food,condition=condition,review_status='Suggested',source_dataset=iid,
                    split=split,group_id=iid,capture_group=iid,pixel_hash=item['quality']['pixel_hash'],sha256=row['sha256'],
                    bbox_status='Pending',bbox_json='[]',label_origin='local_metadata_suggested' if source_folder else 'google_query_suggested'))
        with (stage_path/'metadata.csv').open('w',encoding='utf8',newline='') as stream:
            writer=csv.DictWriter(stream,fieldnames=list(metadata[0]))
            writer.writeheader()
            writer.writerows(metadata)
        write_json(stage_path/'sources.json',sources)
        write_json(stage_path/'manifest.json',dict(classes=sorted(food+'_'+s for s in stages),
            label_mode='google_query_suggested',public=False,food_name=food_name,
            catalog_sha256=hashlib.sha256((stage_path/'metadata.csv').read_bytes()).hexdigest()))
        write_json(stage_path/'filter-report.json',report)
        shutil.move(str(stage_path),str(dest))
    return str(dest)
