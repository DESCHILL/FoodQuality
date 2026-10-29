from __future__ import annotations

import csv
import fnmatch
import hashlib
import json
import math
import shutil
import sqlite3
import tempfile
import urllib.request
import zipfile
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path

import numpy as np
import requests
from PIL import Image, ImageOps
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

FOODS = dict(beef="meat", pork="meat", chicken="poultry", fish="seafood",
             apple="fruit", banana="fruit", tomato="vegetable", potato="vegetable")
STATES = {f: (["fresh_raw", "visibly_spoiled", "cooked"] if f in list(FOODS)[:4]
              else ["fresh", "visibly_spoiled"]) for f in FOODS}
CLASSES = [f + "_" + s for f in FOODS for s in STATES[f]]
EXTENSIONS = {".jpg", ".jpeg", ".png", ".webp", ".bmp", ".tif", ".tiff"}


def digest(data):
    return hashlib.sha256(data).hexdigest()


def phash(image):
    a = np.asarray(image.convert("L").resize((32, 32)), dtype=float)
    basis = np.cos(np.pi / 32 * (np.arange(32) + .5)[None, :] * np.arange(8)[:, None])
    low = (basis @ a @ basis.T).flatten()
    bits = low > np.median(low[1:])
    return f"{sum(int(b) << i for i, b in enumerate(bits)):016x}"


def write_csv(path, rows, fields=None):
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8-sig") as stream:
        writer = csv.DictWriter(stream, fieldnames=fields or list(rows[0]) if rows else fields or ["image_id"])
        writer.writeheader()
        writer.writerows(rows)


def valid_label(food, condition):
    return food in STATES and condition in STATES[food]


class Dataset:
    def __init__(self, root="workspace"):
        self.root = Path(root).resolve()
        for folder in ("config", "data/raw", "data/quarantine", "metadata", "reports", "releases"):
            (self.root / folder).mkdir(parents=True, exist_ok=True)
        self.db = sqlite3.connect(self.root / "metadata/catalog.sqlite")
        self.db.row_factory = sqlite3.Row
        self.db.executescript('''
        CREATE TABLE IF NOT EXISTS images (
          image_id TEXT PRIMARY KEY, source_dataset TEXT, original_path TEXT,
          raw_path TEXT, sha256 TEXT, pixel_hash TEXT, phash TEXT,
          food_name TEXT, category TEXT, condition TEXT,
          review_status TEXT DEFAULT 'Pending', notes TEXT DEFAULT '',
          group_id TEXT, capture_group TEXT, split TEXT DEFAULT '',
          error TEXT DEFAULT '', duplicate_of TEXT DEFAULT '',
          bbox_json TEXT DEFAULT '[]', bbox_status TEXT DEFAULT 'Pending');
        CREATE TABLE IF NOT EXISTS sources (id TEXT PRIMARY KEY, config TEXT);
        CREATE TABLE IF NOT EXISTS audit (time TEXT, action TEXT, detail TEXT);
        CREATE TABLE IF NOT EXISTS settings (key TEXT PRIMARY KEY, value TEXT);
        ''')
        saved = self.db.execute("SELECT value FROM settings WHERE key='taxonomy'").fetchone()
        self.taxonomy = json.loads(saved[0]) if saved else None
        self._load_taxonomy()
        existing_columns = {r[1] for r in self.db.execute('PRAGMA table_info(images)')}
        for column in ('original_label', 'reviewer', 'label_evidence'):
            if column not in existing_columns:
                self.db.execute(f"ALTER TABLE images ADD COLUMN {column} TEXT DEFAULT ''")
        self.db.commit()

    def _load_taxonomy(self):
        self.foods = self.taxonomy['foods'] if self.taxonomy else dict(FOODS)
        self.states = ({f: list(self.taxonomy['conditions'][f] if isinstance(self.taxonomy['conditions'], dict) else self.taxonomy['conditions']) for f in self.foods}
                       if self.taxonomy else dict(STATES))
        self.classes = [f + '_' + s for f in self.foods for s in self.states[f]]

    def valid_label(self, food, condition):
        return food in self.states and condition in self.states[food]

    def set_taxonomy(self, filename):
        taxonomy = json.loads(Path(filename).read_text(encoding='utf-8'))
        if not isinstance(taxonomy.get('foods'), dict) or not taxonomy['foods'] or not taxonomy.get('conditions'):
            raise ValueError('Taxonomy requires foods and conditions')
        conditions = taxonomy['conditions']
        if isinstance(conditions, dict) and set(conditions) != set(taxonomy['foods']):
            raise ValueError('Per-food conditions must cover exactly the configured foods')
        levels = list(conditions.values()) if isinstance(conditions, dict) else [conditions]
        if any(not isinstance(values, list) or not values for values in levels):
            raise ValueError('Each condition list must be nonempty')
        names = list(taxonomy['foods']) + [s for values in levels for s in values]
        if any(not isinstance(n, str) or not n or not all(c.isascii() and (c.isalnum() or c == '_') for c in n) for n in names):
            raise ValueError('Taxonomy ids must use ASCII letters, numbers and underscores')
        if any(len(set(values)) != len(values) for values in levels):
            raise ValueError('Duplicate conditions')
        if self.rows() and taxonomy != self.taxonomy:
            raise ValueError('Use a new workspace to change taxonomy after ingestion')
        self.taxonomy = taxonomy
        self._load_taxonomy()
        with self.db:
            self.db.execute("INSERT OR REPLACE INTO settings VALUES ('taxonomy',?)", (json.dumps(taxonomy),))
            self.audit('taxonomy', taxonomy)
        (self.root / 'config/taxonomy.json').write_text(json.dumps(taxonomy, indent=2, ensure_ascii=False), encoding='utf-8')
        return len(self.classes)

    def close(self):
        self.db.close()

    def audit(self, action, detail):
        self.db.execute("INSERT INTO audit VALUES (?,?,?)", (datetime.now(timezone.utc).isoformat(), action, json.dumps(detail, ensure_ascii=False)))

    def sources(self):
        return {r["id"]: json.loads(r["config"]) for r in self.db.execute("SELECT * FROM sources")}

    def rows(self):
        return [dict(r) for r in self.db.execute("SELECT * FROM images ORDER BY image_id")]

    def configure(self, filename):
        config = json.loads(Path(filename).read_text(encoding="utf-8"))
        sources = config["sources"]
        for s in sources:
            if not s.get("id") or not all(c.isalnum() or c in "_-" for c in s["id"]):
                raise ValueError("Source id must contain letters, numbers, _ or -")
            if s.get("license_status", "Pending") not in ("Approved", "Pending", "Rejected"):
                raise ValueError("Invalid license_status")
            if s.get("license_status") == "Approved" and not all(s.get(k) for k in ("license_id", "license_url", "checked_at", "source_url")):
                raise ValueError("Approved license requires id, URL, source URL and checked_at")
            for label in s.get("mapping", {}).values():
                if not self.valid_label(*label) and not (label[0] in self.foods and label[1] == 'unknown'):
                    raise ValueError(f"Invalid mapping: {label}")
        with self.db:
            for s in sources:
                self.db.execute("INSERT OR REPLACE INTO sources VALUES (?,?)", (s["id"], json.dumps(s)))
            self.db.execute("UPDATE images SET split=''")
            self.audit("configure", sources)

    def download(self, source_id):
        """Download an explicitly configured HTTPS ZIP; reject unsafe archive entries."""
        s = self.sources()[source_id]
        if s.get("license_status") != "Approved":
            raise ValueError("Verify and approve source license before downloading")
        url = s.get("download_url", "")
        if not url.startswith("https://"):
            raise ValueError("Set an HTTPS download_url for a ZIP archive")
        dest = self.root / "data/raw" / source_id / "download"
        if dest.exists():
            raise ValueError("Download already exists; use a new source version/id")
        with tempfile.TemporaryDirectory(dir=self.root / "data/raw") as temp:
            archive = Path(temp) / "source.zip"
            with requests.get(url, stream=True, timeout=(30, 120)) as response, archive.open("wb") as out:
                response.raise_for_status()
                size = 0
                last_report = 0
                for chunk in response.iter_content(1024 * 1024):
                    size += len(chunk)
                    if size > s.get("max_download_bytes", 5_000_000_000):
                        raise ValueError("Download exceeds configured size limit")
                    out.write(chunk)
                    if size - last_report >= 25 * 1024 * 1024:
                        print(f"{source_id}: downloaded {size // (1024 * 1024)} MiB", flush=True)
                        last_report = size
            if s.get("archive_sha256") and digest(archive.read_bytes()) != s["archive_sha256"]:
                raise ValueError("Archive SHA256 mismatch")
            extracted = Path(temp) / "extracted"
            extracted.mkdir()
            with zipfile.ZipFile(archive) as z:
                if sum(i.file_size for i in z.infolist()) > s.get("max_extract_bytes", 15_000_000_000):
                    raise ValueError("Archive exceeds extraction limit")
                for i in z.infolist():
                    target = (extracted / i.filename).resolve()
                    if not target.is_relative_to(extracted.resolve()) or "\\" in i.filename or ":" in i.filename or (i.external_attr >> 16) & 0o170000 == 0o120000:
                        raise ValueError("Unsafe ZIP entry")
                z.extractall(extracted)
            dest.parent.mkdir(parents=True, exist_ok=True)
            shutil.move(str(extracted), str(dest))
        with self.db:
            self.audit("download", {"source": source_id, "url": url})
        return str(dest)

    def ingest(self, source_id, directory):
        s = self.sources()[source_id]
        directory = Path(directory).resolve()
        if not directory.is_dir():
            raise ValueError("Input directory does not exist")
        added = 0
        with self.db:
            for p in sorted(directory.rglob("*")):
                if p.suffix.lower() not in EXTENSIONS or not p.is_file():
                    continue
                relative = p.relative_to(directory).as_posix()
                if any(fnmatch.fnmatch(relative, pattern) for pattern in s.get('exclude_globs', [])):
                    continue
                data = p.read_bytes()
                sha = digest(data)
                iid = digest(f"{source_id}:{relative}:{sha}".encode())[:24]
                if self.db.execute("SELECT 1 FROM images WHERE image_id=?", (iid,)).fetchone():
                    continue
                matches = [s["mapping"][part] for part in p.relative_to(directory).parts[:-1] if part in s.get("mapping", {})]
                food, condition = matches[-1] if matches else ("unknown", "unknown")
                error, pixel, perceptual = "", "", ""
                try:
                    with Image.open(p) as im:
                        im = ImageOps.exif_transpose(im).convert("RGB")
                        im.load()
                        pixel = digest(str(im.size).encode() + im.tobytes())
                        perceptual = phash(im)
                except Exception as exc:
                    error = str(exc)
                bucket = "quarantine" if error or s.get("license_status") != "Approved" else "raw"
                dest = self.root / "data" / bucket / source_id / (iid + p.suffix.lower())
                dest.parent.mkdir(parents=True, exist_ok=True)
                dest.write_bytes(data)
                capture = s.get("groups", {}).get(relative, "")
                self.db.execute('''INSERT INTO images
                  (image_id,source_dataset,original_path,raw_path,sha256,pixel_hash,phash,
                   food_name,category,condition,group_id,capture_group,error)
                   VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)''',
                  (iid, source_id, relative, str(dest.relative_to(self.root)), sha, pixel,
                   perceptual, food, self.foods.get(food, "unknown"), condition, iid, capture, error))
                original_labels = [part for part in p.relative_to(directory).parts[:-1] if part in s.get('mapping', {})]
                self.db.execute('UPDATE images SET original_label=? WHERE image_id=?', ('/'.join(original_labels), iid))
                added += 1
                if added % 250 == 0:
                    print(f'{source_id}: inspected {added} images', flush=True)
            if added:
                self.db.execute("UPDATE images SET split=''")
            self.audit("ingest", {"source": source_id, "added": added})
        self.deduplicate()
        self.snapshot()
        return added

    def deduplicate(self, threshold=None):
        if threshold is None:
            saved = self.db.execute("SELECT value FROM settings WHERE key='phash_threshold'").fetchone()
            threshold = int(saved[0]) if saved else 6
        if not 0 <= threshold <= 64:
            raise ValueError("pHash threshold must be 0..64")
        sources = self.sources()
        rows = sorted(self.rows(), key=lambda r: (sources[r['source_dataset']].get('license_status') != 'Approved', r['review_status'] != 'Approved', r['image_id']))
        parent = {r["image_id"]: r["image_id"] for r in rows}
        def find(x):
            while parent[x] != x:
                parent[x] = parent[parent[x]]
                x = parent[x]
            return x
        def union(a, b):
            a, b = find(a), find(b)
            parent[max(a, b)] = min(a, b)
        report, exact, captures, hashes, duplicates = [], {}, {}, [], {}
        for r in rows:
            iid = r["image_id"]
            key = r["pixel_hash"] or r["sha256"]
            if key in exact:
                other = exact[key]
                union(iid, other)
                duplicates[iid] = other
                report.append(dict(image_id=iid, other_id=other, kind="exact", distance=0))
            else:
                exact[key] = iid
            if r["capture_group"]:
                key = r["capture_group"]
                if key in captures:
                    union(iid, captures[key])
                captures[key] = iid
            if r["phash"]:
                value = int(r["phash"], 16)
                for other, v in hashes:
                    distance = (value ^ v).bit_count()
                    if distance <= threshold:
                        union(iid, other)
                        report.append(dict(image_id=iid, other_id=other, kind="near", distance=distance))
                hashes.append((iid, value))
        # Exact duplicates with conflicting labels must all be reviewed.
        by_pixel = defaultdict(list)
        for r in rows:
            by_pixel[r["pixel_hash"] or r["sha256"]].append(r)
        with self.db:
            self.db.execute("INSERT OR REPLACE INTO settings VALUES ('phash_threshold',?)", (str(threshold),))
            for r in rows:
                self.db.execute("UPDATE images SET group_id=?,duplicate_of=?,split='' WHERE image_id=?", (find(r["image_id"]), duplicates.get(r["image_id"], ""), r["image_id"]))
            for group in by_pixel.values():
                if len({(r["food_name"], r["condition"]) for r in group}) > 1:
                    for r in group:
                        self.db.execute("UPDATE images SET review_status='Pending',notes='Conflicting labels on exact duplicate' WHERE image_id=?", (r["image_id"],))
            self.audit("deduplicate", {"threshold": threshold, "pairs": len(report)})
        write_csv(self.root / "reports/duplicate_report.csv", report, ["image_id", "other_id", "kind", "distance"])
        return len(report)

    def snapshot(self):
        rows, sources = self.rows(), self.sources()
        for r in rows:
            s = sources[r["source_dataset"]]
            r.update({k: s.get(k, "") for k in ("source_url", "license_id", "license_status")})
        write_csv(self.root / "metadata/images.csv", rows)
        write_csv(self.root / "metadata/provenance.csv", [{k: r[k] for k in ("image_id", "source_dataset", "original_path", "sha256")} for r in rows])
        return rows

    def statistics(self):
        counts = Counter((r["food_name"] + "_" + r["condition"], r["split"]) for r in self.rows() if self.eligible(r))
        return [dict(label=c, train=counts[c, "train"], val=counts[c, "val"], test=counts[c, "test"],
                     unsplit=counts[c, ""], total=sum(counts[c, s] for s in ("train", "val", "test", "")),
                     status="LOW_DATA" if sum(counts[c, s] for s in ("train", "val", "test", "")) < 300 else "OK") for c in self.classes]

    def quality_report(self):
        rows = self.rows()
        sources = self.sources()
        changed, mapped = 0, 0
        for r in rows:
            mapping = sources[r['source_dataset']].get('mapping', {})
            labels = [mapping[p] for p in Path(r['original_path']).parts[:-1] if p in mapping]
            if labels:
                mapped += 1
                changed += list(labels[-1]) != [r['food_name'], r['condition']]
        return dict(total=len(rows), eligible=sum(self.eligible(r) for r in rows),
                    corrupt=sum(bool(r['error']) for r in rows),
                    exact_duplicates=sum(bool(r['duplicate_of']) for r in rows),
                    review_status=dict(Counter(r['review_status'] for r in rows)),
                    corrected_vs_current_mapping=changed,
                    correction_fraction=changed / mapped if mapped else None,
                    unused_sources=[sid for sid in sources if not any(r['source_dataset'] == sid and self.eligible(r) for r in rows)])

    def coverage(self):
        rows = self.rows()
        stats = self.statistics()
        candidates = Counter(r['food_name'] for r in rows if not r['error'] and not r['duplicate_of'])
        result = []
        for food in self.foods:
            for condition in self.states[food]:
                stat = next(s for s in stats if s['label'] == food + '_' + condition)
                reviewed = stat['total']
                result.append(dict(food=food, condition=condition,
                                   candidates_food_only=candidates[food],
                                   approved=reviewed, train=stat['train'], val=stat['val'], test=stat['test'],
                                   missing_to_300=max(0, 300-reviewed),
                                   status='MISSING' if not reviewed else 'LOW_DATA' if reviewed < 300 else 'REVIEW_SPLITS'))
        write_csv(self.root / 'reports/coverage.csv', result)
        summary = dict(target_foods=len(self.foods), target_classes=len(self.classes),
                       downloaded_records=len(rows), candidate_foods=len([f for f in self.foods if candidates[f]]),
                       approved_classes=sum(r['approved'] > 0 for r in result),
                       missing_classes=[r['food']+'_'+r['condition'] for r in result if not r['approved']],
                       ready_for_training=all(r['approved'] >= 300 and r['train'] and r['val'] and r['test'] for r in result))
        (self.root / 'reports/readiness.json').write_text(json.dumps(summary, indent=2), encoding='utf-8')
        return summary

    def eligible(self, r):
        return (not r["error"] and not r["duplicate_of"] and r["review_status"] == "Approved"
                and self.valid_label(r["food_name"], r["condition"])
                and self.sources()[r["source_dataset"]].get("license_status") == "Approved")

    def review_export(self, path=None):
        path = Path(path or self.root / "reports/FoodQuality-8.xlsx")
        wb = Workbook()
        wb.remove(wb.active)
        datasets = {"Images": self.snapshot(), "Sources_Licenses": self.sources().values(),
                    "Label_Mapping": [dict(source=s["id"], original=k, food_name=v[0], condition=v[1]) for s in self.sources().values() for k, v in s.get("mapping", {}).items()],
                    "Statistics": self.statistics(),
                    "Target_Labels": [dict(food=f, condition=s) for f in self.foods for s in self.states[f]]}
        for title, data in datasets.items():
            data = list(data)
            ws = wb.create_sheet(title)
            fields = list(dict.fromkeys(k for row in data for k in row)) or ["image_id"]
            ws.append(fields)
            for row in data:
                ws.append([json.dumps(row.get(k), ensure_ascii=False) if isinstance(row.get(k), (dict, list)) else row.get(k, "") for k in fields])
            for cells in ws:
                for cell in cells:
                    if isinstance(cell.value, str):
                        cell.data_type = "s"  # Do not evaluate source text as Excel formula.
            for cell in ws[1]:
                cell.font = Font(color="FFFFFF", bold=True)
                cell.fill = PatternFill("solid", fgColor="174A45")
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            for column in ws.columns:
                ws.column_dimensions[column[0].column_letter].width = min(45, max(16, len(str(column[0].value)) + 3))
            if title == "Images":
                for field in ("review_status", "bbox_status"):
                    if field in fields:
                        dv = DataValidation(type="list", formula1='"Pending,Approved,Rejected"')
                        dv.error = "Choose Pending, Approved or Rejected"
                        dv.showErrorMessage = True
                        ws.add_data_validation(dv)
                        col = ws.cell(1, fields.index(field) + 1).column_letter
                        dv.add(f"{col}2:{col}{max(2, ws.max_row)}")
        path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(path)
        return str(path)

    def validate_boxes(self, boxes):
        if not isinstance(boxes, list):
            raise ValueError("bbox_json must be a JSON list")
        for b in boxes:
            if not isinstance(b, dict) or b.get("food_name") not in self.foods:
                raise ValueError("Invalid box food_name")
            vals = [b.get(k) for k in ("x", "y", "w", "h")]
            if not all(isinstance(v, (int, float)) and math.isfinite(v) for v in vals):
                raise ValueError("Box coordinates must be finite numbers")
            x, y, w, h = vals
            if not (w > 0 and h > 0 and x-w/2 >= 0 and x+w/2 <= 1 and y-h/2 >= 0 and y+h/2 <= 1):
                raise ValueError("Boxes must be normalized YOLO xywh within image bounds")

    def review_import(self, path):
        wb = load_workbook(path, read_only=True, data_only=False)
        try:
            values = iter(wb["Images"].values)
            headers = next(values)
            current = {r["image_id"]: r for r in self.rows()}
            updates, seen = [], set()
            for values_row in values:
                r = dict(zip(headers, values_row))
                iid = r.get("image_id")
                if iid not in current or iid in seen:
                    raise ValueError(f"Unknown or duplicate image_id: {iid}")
                seen.add(iid)
                for field in ("review_status", "bbox_status"):
                    if r.get(field) not in ("Approved", "Pending", "Rejected"):
                        raise ValueError(f"Invalid {field}: {iid}")
                food, state = r.get("food_name"), r.get("condition")
                if not self.valid_label(food, state) and not (food in [*self.foods, 'unknown'] and state == 'unknown'):
                    raise ValueError(f"Invalid label: {iid}")
                if r["review_status"] == "Approved" and (current[iid]["error"] or not self.valid_label(food, state)):
                    raise ValueError(f"Cannot approve corrupt or unknown image: {iid}")
                if self.taxonomy and r['review_status'] == 'Approved' and not (r.get('reviewer') and r.get('label_evidence')):
                    raise ValueError(f'Custom taxonomy review requires reviewer and label_evidence: {iid}')
                boxes = json.loads(r.get("bbox_json") or "[]")
                self.validate_boxes(boxes)
                updates.append((food, self.foods.get(food, "unknown"), state, r["review_status"], r.get("notes") or "", r.get("capture_group") or "", json.dumps(boxes), r["bbox_status"], r.get('reviewer') or '', r.get('label_evidence') or '', iid))
            with self.db:
                self.db.executemany("UPDATE images SET food_name=?,category=?,condition=?,review_status=?,notes=?,capture_group=?,bbox_json=?,bbox_status=?,reviewer=?,label_evidence=?,split='' WHERE image_id=?", updates)
                self.audit("review_import", {"file": str(path), "count": len(updates)})
        finally:
            wb.close()
        self.deduplicate()
        self.snapshot()
        return len(updates)

    def split(self, seed=42, ratios=(.7, .15, .15)):
        if len(ratios) != 3 or any(x <= 0 for x in ratios) or not math.isclose(sum(ratios), 1):
            raise ValueError("Three positive split ratios must sum to 1")
        groups = defaultdict(list)
        for r in self.rows():
            if self.eligible(r):
                groups[r["group_id"]].append(r)
        totals = Counter(r["food_name"] + "_" + r["condition"] for rs in groups.values() for r in rs)
        assigned = [Counter(), Counter(), Counter()]
        names = ("train", "val", "test")
        with self.db:
            self.db.execute("UPDATE images SET split=''")
            for gid, rs in sorted(groups.items(), key=lambda pair: (-len(pair[1]), digest(f"{seed}:{pair[0]}".encode()))):
                counts = Counter(r["food_name"] + "_" + r["condition"] for r in rs)
                costs = [sum(((assigned[i][c]+n-totals[c]*ratios[i])**2 - (assigned[i][c]-totals[c]*ratios[i])**2) / max(totals[c]*ratios[i], 1) for c, n in counts.items()) for i in range(3)]
                choice = min(range(3), key=lambda i: costs[i])
                assigned[choice].update(counts)
                for r in rs:
                    self.db.execute("UPDATE images SET split=? WHERE image_id=?", (names[choice], r["image_id"]))
            self.audit("split", {"seed": seed, "ratios": ratios})
        self.snapshot()
        write_csv(self.root / "reports/dataset_statistics.csv", self.statistics())
        return {names[i]: sum(assigned[i].values()) for i in range(3)}

    def export(self, version, public=False, selected_classes=None):
        if not version or not all(c.isalnum() or c in "._-" for c in version) or version in (".", ".."):
            raise ValueError("Invalid release version")
        dest = self.root / "releases" / version
        if dest.exists():
            raise ValueError("Release exists; choose a new version")
        classes = list(self.classes) if selected_classes is None else list(dict.fromkeys(selected_classes))
        if not classes or any(c not in self.classes for c in classes):
            raise ValueError('Select valid nonempty classes for the release')
        export_foods = list(self.foods) if selected_classes is None else [f for f in self.foods if any(f+'_'+s in classes for s in self.states[f])]
        if self.taxonomy or selected_classes is not None:
            stats = self.statistics()
            missing = [s['label'] for s in stats if s['label'] in classes and not all(s[k] for k in ('train', 'val', 'test'))]
            if missing:
                raise ValueError(f'Incomplete target dataset: {len(missing)} classes lack approved train/val/test images. Run coverage.')
        rows = [r for r in self.rows() if self.eligible(r) and r['food_name']+'_'+r['condition'] in classes]
        if not rows or any(r["split"] not in ("train", "val", "test") for r in rows):
            raise ValueError("No approved data, or split is missing; run review and split")
        seen = {}
        sources = self.sources()
        for r in rows:
            for key in (r["group_id"], r["sha256"], r["pixel_hash"]):
                if key and key in seen and seen[key] != r["split"]:
                    raise ValueError("Cross-split data leakage")
                seen[key] = r["split"]
            if public and not sources[r["source_dataset"]].get("redistribution_allowed", False):
                raise ValueError("Public export requires redistribution_allowed for every source")
            if digest((self.root / r["raw_path"]).read_bytes()) != r["sha256"]:
                raise ValueError("Raw image changed: " + r["image_id"])
            self.validate_boxes(json.loads(r["bbox_json"]))
            boxes = json.loads(r['bbox_json'])
            if boxes and r['bbox_status']=='Approved':
                if len(boxes)!=1 or boxes[0]['food_name']!=r['food_name']:
                    raise ValueError('Classification release needs one reviewed food crop per image; split multi-food images into reviewed crops first')
            r['classification_crop_json'] = json.dumps(boxes[0] if boxes and r['bbox_status']=='Approved' else None)
        with tempfile.TemporaryDirectory(dir=self.root / "releases") as temp:
            stage = Path(temp) / "release"
            stage.mkdir()
            for split in ("train", "val", "test"):
                for label in classes:
                    (stage / "classification" / split / label).mkdir(parents=True)
                for kind in ("images", "labels"):
                    (stage / "detection" / kind / split).mkdir(parents=True)
            detection_count = 0
            for r in rows:
                label = r["food_name"] + "_" + r["condition"]
                target = stage / "classification" / r["split"] / label / (r["image_id"] + ".jpg")
                boxes = json.loads(r["bbox_json"])
                with Image.open(self.root / r["raw_path"]) as im:
                    full = ImageOps.exif_transpose(im).convert("RGB")
                    crop = full
                    if boxes and r['bbox_status']=='Approved':
                        b=boxes[0]
                        left=max(0,int((b['x']-b['w']/2)*full.width))
                        top=max(0,int((b['y']-b['h']/2)*full.height))
                        right=min(full.width,max(left+1,int((b['x']+b['w']/2)*full.width)))
                        bottom=min(full.height,max(top+1,int((b['y']+b['h']/2)*full.height)))
                        crop=full.crop((left,top,right,bottom))
                    crop.save(target,quality=95)
                    if boxes and r['bbox_status']=='Approved':
                        full.save(stage/'detection/images'/r['split']/target.name,quality=95)
                if boxes and r["bbox_status"] == "Approved":
                    lines = [f"{export_foods.index(b['food_name'])} {b['x']} {b['y']} {b['w']} {b['h']}" for b in boxes]
                    (stage / "detection/labels" / r["split"] / (r["image_id"] + ".txt")).write_text("\n".join(lines), encoding="utf-8")
                    detection_count += 1
            (stage / "detection/data.yaml").write_text("path: " + json.dumps(str(dest / "detection")) + "\ntrain: images/train\nval: images/val\ntest: images/test\nnames: " + json.dumps(export_foods) + "\n", encoding="utf-8")
            write_csv(stage / "metadata.csv", rows)
            for split in ("train", "val", "test"):
                write_csv(stage / "splits" / (split + ".csv"), [r for r in rows if r["split"] == split])
            write_csv(stage / "statistics.csv", [s for s in self.statistics() if s['label'] in classes])
            (stage / "sources.json").write_text(json.dumps(sources, indent=2, ensure_ascii=False), encoding="utf-8")
            (stage / "quality_report.json").write_text(json.dumps(self.quality_report(), indent=2), encoding="utf-8")
            manifest = dict(version=version, created_at=datetime.now(timezone.utc).isoformat(), images=len(rows), detection_images=detection_count,
                            classes=classes, public=public, scope='full' if selected_classes is None else 'explicit_subset',
                            full_target_classes=self.classes, missing_classes=[s["label"] for s in self.statistics() if s['label'] in classes and not s["total"]],
                            catalog_sha256=digest((stage / "metadata.csv").read_bytes()))
            (stage / "manifest.json").write_text(json.dumps(manifest, indent=2), encoding="utf-8")
            shutil.move(str(stage), str(dest))
        with self.db:
            self.audit("export", manifest)
        return str(dest)
