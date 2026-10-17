"""Synthetic fixtures only; no food images or training claims."""
import json
from pathlib import Path
import numpy as np
from PIL import Image
from openpyxl import load_workbook
from foodquality8 import Dataset


def main():
    root = Path("demo_workspace")
    d = Dataset(root)
    config = {"sources": [dict(id="synthetic", source_url="local:synthetic-demo", license_id="OWN",
        license_url="local:generated-by-demo", checked_at="2026-10-16", license_status="Approved",
        owner="Demo generator", redistribution_allowed=True,
        mapping={"FreshApple": ["apple", "fresh"], "RottenApple": ["apple", "visibly_spoiled"]})]}
    config_path = root / "config/sources.json"
    config_path.write_text(json.dumps(config, indent=2))
    d.configure(config_path)
    rng = np.random.default_rng(2026)
    for label in config["sources"][0]["mapping"]:
        folder = root / "synthetic_input" / label
        folder.mkdir(parents=True, exist_ok=True)
        for i in range(20):
            Image.fromarray(rng.integers(0, 256, (80, 80, 3), dtype=np.uint8)).save(folder / f"{i}.png")
    d.ingest("synthetic", root / "synthetic_input")
    review = d.review_export()
    wb = load_workbook(review)
    ws = wb["Images"]
    col = [c.value for c in ws[1]].index("review_status") + 1
    for row in range(2, ws.max_row+1):
        ws.cell(row, col, "Approved")  # Synthetic fixture, never auto-approve real data.
    wb.save(review)
    wb.close()
    d.review_import(review)
    print(d.split())
    version = "demo-v" + str(len(list((d.root / "releases").iterdir())) + 1)
    print(d.export(version))
    print(d.review_export())
    d.close()


if __name__ == "__main__":
    main()
